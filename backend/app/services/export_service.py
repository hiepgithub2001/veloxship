"""Excel export for bills (xlsx via openpyxl)."""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

STATUS_LABELS = {
    "created": "Đã tạo",
    "picked_up": "Đã lấy hàng",
    "in_transit": "Đang vận chuyển",
    "delivered": "Đã giao",
    "returned": "Hoàn trả",
    "cancelled": "Đã hủy",
}

_HEADERS = [
    "Mã vận đơn",
    "Người gửi",
    "SĐT người gửi",
    "Người nhận",
    "SĐT người nhận",
    "Trạng thái",
    "Tổng cước (VND)",
    "Ngày tạo",
]


def _party_name(bill, side: str) -> str:
    """Historical name from the immutable snapshot, falling back to the customer."""
    snapshot = bill.sender_snapshot if side == "sender" else bill.receiver_snapshot
    if snapshot and snapshot.get("name"):
        return snapshot["name"]
    customer = bill.sender if side == "sender" else bill.receiver
    return customer.name if customer else ""


def _party_phone(bill, side: str) -> str:
    snapshot = bill.sender_snapshot if side == "sender" else bill.receiver_snapshot
    if snapshot and snapshot.get("phone"):
        return snapshot["phone"]
    customer = bill.sender if side == "sender" else bill.receiver
    return customer.phone if customer and customer.phone else ""


def _format_datetime(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.strftime("%d/%m/%Y %H:%M")


def build_bills_xlsx(bills) -> io.BytesIO:
    """Build an xlsx workbook of the given bills and return it as a byte buffer."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Phiếu gửi"

    sheet.append(_HEADERS)
    header_font = Font(bold=True)
    for cell in sheet[1]:
        cell.font = header_font

    for bill in bills:
        sheet.append(
            [
                bill.tracking_number,
                _party_name(bill, "sender"),
                _party_phone(bill, "sender"),
                _party_name(bill, "receiver"),
                _party_phone(bill, "receiver"),
                STATUS_LABELS.get(bill.status, bill.status),
                int(bill.fee_total),
                _format_datetime(bill.created_at),
            ]
        )

    # Reasonable column widths for the Vietnamese labels.
    widths = [22, 28, 16, 28, 16, 18, 18, 18]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
